"""
Script to:
1. Connect to PostgreSQL (rv_pbpkghvg) using radiology_ai.db credentials.
2. Initialize the 'radiology_scan' table if not exists.
3. Fetch currently admitted patients from admissions table (discharge_status = 'Admitted').
4. Randomly map each unique scan case in Patient_id.xlsx to a currently admitted patient_id.
5. Update Patient_id.xlsx with patient_id, patient_code, original_patient_id, image, scan_report.
6. Insert all records into PostgreSQL table 'radiology_scan'.
"""
import sys
import os
import shutil
import random
from pathlib import Path

# Add backend to python path
BACKEND_DIR = Path(__file__).resolve().parent.parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import openpyxl
from radiology_ai.db import (
    get_connection,
    init_radiology_scan_table,
    get_currently_admitted_patients,
    insert_scans,
)

EXCEL_PATH = BACKEND_DIR / "Patient_id.xlsx"
BACKUP_PATH = BACKEND_DIR / "Patient_id.xlsx.bak"


def run():
    print(f"Loading Excel file from: {EXCEL_PATH}")
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"Excel file not found at {EXCEL_PATH}")

    # Create backup if not already present
    if not BACKUP_PATH.exists():
        shutil.copy2(EXCEL_PATH, BACKUP_PATH)
        print(f"Created backup at {BACKUP_PATH}")

    # Step 1: Initialize DB table
    print("Initializing PostgreSQL radiology_scan table...")
    init_radiology_scan_table()

    # Step 2: Fetch currently admitted patients
    admitted = get_currently_admitted_patients()
    print(f"Retrieved {len(admitted)} currently admitted patients from PostgreSQL.")
    if not admitted:
        raise ValueError("No currently admitted patients found in database!")

    # Step 3: Read existing Excel data
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    data_rows = rows[1:]
    print(f"Excel has {len(data_rows)} data rows. Original headers: {header}")

    # Find unique original patientIds (UUIDs)
    orig_uuids = []
    for r in data_rows:
        uid = str(r[0]).strip() if r[0] else ""
        if uid and uid not in orig_uuids:
            orig_uuids.append(uid)

    print(f"Found {len(orig_uuids)} unique scan UUIDs across {len(data_rows)} rows.")

    # Randomly assign an admitted patient to each unique UUID
    # Seed with fixed number for reproducibility or random
    random.seed(42)
    # Sample or choice with replacement if needed
    if len(admitted) >= len(orig_uuids):
        assigned_patients = random.sample(admitted, len(orig_uuids))
    else:
        assigned_patients = [random.choice(admitted) for _ in orig_uuids]

    uuid_to_patient = {
        orig_uuids[i]: assigned_patients[i]
        for i in range(len(orig_uuids))
    }

    # Step 4: Prepare updated records and write back to Excel
    new_wb = openpyxl.Workbook()
    new_ws = new_wb.active
    new_ws.title = "RadiologyScans"

    new_headers = [
        "patient_id",
        "patient_code",
        "original_patient_id",
        "x",
        "y",
        "width",
        "height",
        "Target",
        "image",
        "scan_report"
    ]
    new_ws.append(new_headers)

    scan_records_for_db = []

    for r in data_rows:
        orig_uid = str(r[0]).strip() if r[0] else ""
        pat_info = uuid_to_patient.get(orig_uid, {})
        pat_id = pat_info.get("patient_id")
        pat_code = pat_info.get("patient_code")

        x = float(r[1]) if r[1] is not None else None
        y = float(r[2]) if r[2] is not None else None
        w = float(r[3]) if r[3] is not None else None
        h = float(r[4]) if r[4] is not None else None
        target = int(r[5]) if r[5] is not None else 0

        # Empty image and scan report columns
        image = None
        scan_report = None

        new_row = [
            pat_id,
            pat_code,
            orig_uid,
            x,
            y,
            w,
            h,
            target,
            image,
            scan_report
        ]
        new_ws.append(new_row)

        scan_records_for_db.append({
            "patient_id": pat_id,
            "patient_code": pat_code,
            "original_patient_id": orig_uid,
            "x": x,
            "y": y,
            "width": w,
            "height": h,
            "target": target,
            "image": image,
            "scan_report": scan_report
        })

    new_wb.save(EXCEL_PATH)
    print(f"Successfully updated Excel file at {EXCEL_PATH}")

    # Step 5: Store in PostgreSQL radiology_scan table
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            # Check existing count in radiology_scan
            cur.execute("SELECT count(*) FROM radiology_scan;")
            existing_count = cur.fetchone()[0]
            if existing_count > 0:
                print(f"Table radiology_scan currently has {existing_count} rows. Clearing before fresh ingest...")
                cur.execute("TRUNCATE TABLE radiology_scan RESTART IDENTITY;")
                conn.commit()

        inserted_count = insert_scans(scan_records_for_db)
        print(f"Successfully inserted {inserted_count} records into PostgreSQL 'radiology_scan' table!")

        # Verify insertion
        with conn.cursor() as cur:
            cur.execute("""
                SELECT 
                    count(*) as total_rows,
                    count(DISTINCT patient_id) as unique_patients,
                    sum(CASE WHEN target = 1 THEN 1 ELSE 0 END) as positive_targets,
                    sum(CASE WHEN target = 0 THEN 1 ELSE 0 END) as negative_targets
                FROM radiology_scan;
            """)
            summary = cur.fetchone()
            print(f"Verification Summary: Total={summary[0]}, Unique Patients={summary[1]}, Target=1: {summary[2]}, Target=0: {summary[3]}")

    finally:
        conn.close()

    print("Complete!")


if __name__ == "__main__":
    run()
