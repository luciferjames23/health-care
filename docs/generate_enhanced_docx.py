import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        
    # Colors
    NAVY = RGBColor(15, 44, 89)       # #0F2C59 - Main Titles
    TEAL = RGBColor(14, 116, 144)     # #0E7490 - Subtitles & Headers
    SLATE = RGBColor(51, 65, 85)      # #334155 - Body & Subheaders
    DARK_BLUE = RGBColor(30, 58, 138) # #1E3A8A - Section Headings
    CHARCOAL = RGBColor(31, 41, 55)   # #1F2937 - Main Text
    MUTED = RGBColor(100, 116, 139)   # #64748B - Captions & Notes
    
    # Base styling
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(11)
    normal_font.color.rgb = CHARCOAL
    
    def add_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = TEAL
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = DARK_BLUE
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = TEAL
        return p

    def add_p(text, bold_prefix="", italic=False, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = 'Calibri'
            r_bold.font.bold = True
            r_bold.font.color.rgb = SLATE
        if text:
            r_text = p.add_run(text)
            r_text.font.name = 'Calibri'
            r_text.font.italic = italic
            r_text.font.color.rgb = CHARCOAL
        return p

    def add_bullet(bold_prefix, text, space_after=3):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = 'Calibri'
            r_bold.font.bold = True
            r_bold.font.color.rgb = SLATE
        if text:
            r_text = p.add_run(text)
            r_text.font.name = 'Calibri'
            r_text.font.color.rgb = CHARCOAL
        return p

    def add_callout(text, prefix="NOTE: "):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(6.7)
        
        # Style callout box
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
        tcPr.append(shd)
        
        # Left border only (teal accent)
        tcBorders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="0E7490"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(tcBorders)
        
        # Margins
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="120" w:type="dxa"/>
                <w:bottom w:w="120" w:type="dxa"/>
                <w:left w:w="180" w:type="dxa"/>
                <w:right w:w="180" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        if prefix:
            r_pref = p.add_run(prefix)
            r_pref.font.name = 'Calibri'
            r_pref.font.bold = True
            r_pref.font.size = Pt(10)
            r_pref.font.color.rgb = TEAL
            
        r_text = p.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = SLATE
        
        # Spacer
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # DOCUMENT HEADER
    # -------------------------------------------------------------
    add_title("HEALTHCARE PLATFORM · ENTERPRISE AI ARCHITECTURE")
    
    # Document Main Title
    p_main = doc.add_paragraph()
    p_main.paragraph_format.space_before = Pt(2)
    p_main.paragraph_format.space_after = Pt(6)
    r_main = p_main.add_run("Eight Hospital AI Agents — Purpose, Usefulness, Inputs, Outputs & LLM Architecture")
    r_main.font.name = 'Calibri'
    r_main.font.size = Pt(22)
    r_main.font.bold = True
    r_main.font.color.rgb = NAVY
    
    add_callout(
        "Internal reference compiled from live backend services, Agent Studio configurations, and UI workspaces. "
        "Updated Date: October 2026. This platform is not an autonomous conversational chatbot catalogue — the majority of "
        "agents operate as silent, event-driven clinical and financial workflow automation engines requiring mandatory "
        "human-in-the-loop review and digital sign-off.",
        prefix="SYSTEM NOTICE: "
    )

    # -------------------------------------------------------------
    # SECTION 1: QUICK MAP OF THE EIGHT AGENTS
    # -------------------------------------------------------------
    add_h1("1. Quick Map of the Eight Agents")
    
    add_p(
        "You asked for these eight. Official IDs in the product are listed below. ‘Billing revenue pace’ is documented "
        "as AG-08 Billing Transparency Agent — the live billing/revenue explainer tied to estimate vs actual charge pace. "
        "Financial Revenue screens consume related finance APIs but AG-08 is the agent."
    )
    
    agents_summary = [
        ("AG-19: Discharge Summary Agent", 
         "Medical Records & Inpatient Care",
         "Doctor reviews and electronically signs AI-drafted summary; verifies vital stability score & billing clearance.",
         "Groq openai/gpt-oss-120b (synthesis) + openai/gpt-oss-20b (vital stability)",
         "Mandatory Physician Electronic Sign-Off"),

        ("AG-18: Nursing Shift Handover Agent", 
         "Nursing Operations & Ward Care",
         "Registered Nurse (RN) reviews auto-compiled SBAR (Situation, Background, Assessment, Recommendation) bedside card.",
         "Groq openai/gpt-oss-120b (or Meta-Llama-3.3-70B-Instruct)",
         "Mandatory Bedside RN Joint Acknowledgment"),

        ("AG-04: Employee Self-Service & HR Assistant", 
         "Human Resources (HR) & Workforce Operations",
         "Conversational staff assistant for shift rosters, leave balance inquiries, overtime policy & comp-off requests.",
         "Groq openai/gpt-oss-120b (fallback Gemini 2.0 Flash / 3.5 Flash-Lite)",
         "Automated Supervisor Routing for Approvals"),

        ("AG-14: Executive Hospital Operations & Analytics Agent", 
         "Hospital Executive Leadership & Finance",
         "Turns live SQL database metrics into real-time KPI cards and natural-language daily operational briefings.",
         "PostgreSQL SQL Engine + Groq openai/gpt-oss-120b Narrative Overlay",
         "Read-Only Boundary (Strict Zero-PHI Export)"),

        ("AG-15: Inpatient Bed Census & Ward Capacity Forecasting Agent", 
         "Hospital Operations & Bed Management",
         "Projects 7-day inpatient census, ward occupancy, and capacity surge warnings across ICU, CCU, and wards.",
         "Census Predictor (LightGBM + Prophet math) + Groq openai/gpt-oss-120b Overlay",
         "COO / Nursing Supervisor Approval for Bed Moves"),

        ("AG-07: Cashless Insurance Pre-Authorization Dossier Agent", 
         "Insurance Desk & Revenue Cycle Management (RCM)",
         "Synthesizes cashless pre-authorization dossier, compiles clinical evidence & evaluates AI denial-risk score.",
         "Groq openai/gpt-oss-120b + Policy & Clinical Validation Engine (preauth-denial v0.9)",
         "Mandatory Insurance Officer 1-Click Submission"),

        ("AG-20: Insurance Claim Shortfall & Denial Appeal Agent", 
         "Insurance Desk & Revenue Cycle Lead",
         "Identifies denial codes, retrieves EMR clinical proof, and auto-drafts formal reconsideration appeal letters.",
         "Grounded Clinical Evidence Retrieval + Groq openai/gpt-oss-120b Synthesis",
         "Mandatory Insurance Officer Review & Submit"),

        ("AG-08: Billing Transparency & Estimate Variance Explainer Agent", 
         "Hospital Finance & Billing Desk",
         "Audits running bill vs estimate variance and produces plain-language bilingual (English/Tamil) breakdown cards.",
         "Groq openai/gpt-oss-120b Bilingual Explainer",
         "Optional Cashier Print / Digital Sign-Off")
    ]
    
    for agent_title, domain, desc, llm, gate in agents_summary:
        add_h3(f"• {agent_title} ({domain})")
        add_bullet("Operational Workflow: ", desc)
        add_bullet("Primary LLM / Engine: ", llm)
        add_bullet("Human Governance Gate: ", gate, space_after=6)

    # -------------------------------------------------------------
    # SECTION 2: SHARED LLM PLATFORM & ARCHITECTURAL GOVERNANCE
    # -------------------------------------------------------------
    add_h1("2. Shared LLM Platform & Architectural Governance")
    
    add_p(
        "All clinical, administrative, and financial AI agents within the platform operate on a standardized, high-performance, "
        "and HIPAA-compliant inference infrastructure:"
    )
    
    add_bullet(
        "Primary LLM Inference Engine: ",
        "Most clinical and financial workflow agents call Groq’s OpenAI-compatible Chat Completions API (Application Programming Interface) "
        "powered by ultra-low-latency LPU (Language Processing Unit) hardware using the production model openai/gpt-oss-120b."
    )
    add_bullet(
        "Multi-Tier Resilient Fallbacks: ",
        "To ensure continuous high-availability, agents seamlessly fall back to Meta-Llama-3.3-70B-Instruct on Groq, "
        "Google Gemini (gemini-2.0-flash / gemini-3.5-flash-lite via LLM_API_KEY), or Databricks Meta-Llama-3.3-70B-Instruct."
    )
    add_bullet(
        "Specialized Clinical Sub-Models: ",
        "Dedicated tasks utilize specialized model configurations, such as the lightweight openai/gpt-oss-20b on Groq "
        "for vital-sign clinical stability scoring."
    )
    add_bullet(
        "Deterministic Grounding First: ",
        "Hospital operational statistics, bed counts, billing ledgers, and financial KPIs (Key Performance Indicators) "
        "are strictly computed in PostgreSQL database engines first using deterministic SQL (Structured Query Language). "
        "The LLM (Large Language Model) is strictly utilized to translate deterministic data into natural-language briefings and summaries."
    )
    add_bullet(
        "Mandatory Human-in-the-Loop Governance: ",
        "No AI agent autonomously signs a clinical discharge, submits an insurance TPA (Third-Party Administrator) pre-authorization pack, "
        "dispatches an appeal letter, or alters billing tariffs without explicit human confirmation and digital authorization."
    )

    # -------------------------------------------------------------
    # SECTION 3: AGENT-BY-AGENT DETAILED SPECIFICATION
    # -------------------------------------------------------------
    add_h1("3. Comprehensive Agent-by-Agent Specifications")
    
    # -------------------------------------------------------------
    # AG-19
    # -------------------------------------------------------------
    add_h2("AG-19 · Discharge Summary Agent")
    add_p("Medical Records · High Clinical Tier · Published · v1.1.0", bold_prefix="Domain & Governance: ", italic=True)
    
    add_p(
        "When an attending clinician marks an inpatient as likely for discharge, this agent automatically compiles the entire "
        "multidisciplinary inpatient stay into an NABH (National Accreditation Board for Hospitals & Healthcare Providers) compliant "
        "structured discharge summary draft. It concurrently verifies financial billing clearance and vital-sign stability before "
        "the patient is presented as ready for discharge.",
        bold_prefix="Purpose: "
    )
    
    add_p(
        "Eliminates the tedious 45–60 minute manual documentation burden on consultant physicians. Doctors review a pre-compiled, "
        "clinically grounded draft, make edits if required, and electronically sign. This shortens discharge TAT (Turnaround Time), "
        "prevents missed medication reconciliations, ensures 100% NABH compliance, and delivers bilingual home-care instructions to the family.",
        bold_prefix="Why It Is Useful: "
    )
    
    add_h3("Inputs Given to the Agent:")
    add_bullet("Patient Encounter & Demographics: ", "Patient ID, Admission ID, admission date/time, ward/bed history, attending physician, and clinical specialty.")
    add_bullet("Diagnoses & Clinical Coding: ", "Primary provisional and final diagnoses mapped with ICD-10 (International Classification of Diseases, 10th Revision) codes.")
    add_bullet("Vital Signs Stream: ", "Complete longitudinal vitals record (Blood Pressure, Heart Rate, SpO2, Temperature, Respiratory Rate) from EMR (Electronic Medical Record).")
    add_bullet("Laboratory & Diagnostic Investigations: ", "LIS (Laboratory Information System) lab results, critical value alerts, microbiology cultures, radiology, and ECG reports.")
    add_bullet("Surgical & Procedure Notes: ", "OT (Operation Theatre) operative notes, anesthesia logs, surgical implants used, and procedural findings.")
    add_bullet("Medication Administration History: ", "eMAR (Electronic Medication Administration Record) inpatient administrations and hospital pharmacy formulary records.")
    add_bullet("Financial Clearance Status: ", "Real-time billing settlement confirmation from Hospital Billing API (Application Programming Interface).")
    
    add_h3("Outputs Produced by the Agent:")
    add_bullet("Structured NABH Discharge Summary Draft: ", "Comprehensive multi-section clinical document adhering to NABH Clinical Documentation Standards v5.0.")
    add_bullet("Vital Stability Clearance & Readiness Score: ", "AI-evaluated vital-sign stability index confirming clinical readiness for safe home transfer.")
    add_bullet("Reconciled Discharge Medication Schedule: ", "Clear medication table specifying drug generic name, dosage, frequency, route, duration, and meal relation.")
    add_bullet("Bilingual Patient Home-Care Instructions: ", "Plain-language home recovery guidelines, wound care, and activity restrictions in English and Tamil.")
    add_bullet("Emergency Red-Flag Warnings & Follow-Up Plan: ", "Critical warning signs requiring immediate ED (Emergency Department) return and scheduled OPD (Outpatient Department) follow-up appointment.")
    add_bullet("Persisted Audit Record: ", "Validated draft saved into gold database table generated-discharge-summaries awaiting mandatory physician digital signature.")
    
    add_h3("Technical & Operational Details:")
    add_bullet("How Staff Use It: ", "Doctor Workspace / Discharge Command Centre. Physician clicks 'Sign & Finalize' after reviewing. Zero autonomous discharge.")
    add_bullet("Delivery Mode & Primary Users: ", "Clinical Workflow Automation (Doctor Workspace, Patient 360). Primary users: Consultant Doctors, Medical Records.")
    add_bullet("LLM Models: ", "Primary synthesis: openai/gpt-oss-120b on Groq LPU (Language Processing Unit). Vital gate: openai/gpt-oss-20b. Fallback: Gemini 3.5 Flash-Lite / Meta-Llama-3.3-70B. Temp: 0.10–0.30.")
    add_bullet("Tools & Integrations: ", "EMR Gateway, LIS Labs API, Pharmacy Formulary API, Hospital Billing DB, Document Generator.")
    add_bullet("Knowledge & SOPs: ", "NABH Clinical Documentation v5.0, Inpatient Discharge SOP v3.1, High-Alert Medications v4.0.")
    add_bullet("Quality Benchmarks: ", "Studio Eval EV-8925: 115 test cases, 100% completeness, 99.4% groundedness, 0.1% hallucination, 1.18s median latency.", space_after=12)

    # -------------------------------------------------------------
    # AG-18
    # -------------------------------------------------------------
    add_h2("AG-18 · Nursing Shift Handover Agent")
    add_p("Nursing Operations · Medium Clinical Summarisation Tier · Production-Pilot · v0.8.2", bold_prefix="Domain & Governance: ", italic=True)
    
    add_p(
        "At each nursing shift change (07:00, 15:00, 23:00), the agent reads the previous shift's longitudinal vitals, "
        "eMAR (Electronic Medication Administration Record) administrations, nursing intervention tasks, and physician orders "
        "to pre-draft a standardized bedside SBAR (Situation, Background, Assessment, Recommendation) handover card while highlighting "
        "high-alert medications and deteriorating clinical trends.",
        bold_prefix="Purpose: "
    )
    
    add_p(
        "Replaces disorganized paper handover notes and lengthy verbal handovers that previously took over 10 minutes per bed. "
        "With pre-populated SBAR cards, bedside handover is completed in ~90 seconds with zero omissions of high-alert drug doses, "
        "pending stat lab tests, or rising EWS (Early Warning Score) indications.",
        bold_prefix="Why It Is Useful: "
    )
    
    add_h3("Inputs Given to the Agent:")
    add_bullet("Ward & Patient Identifier: ", "Ward ID, Bed/Room Number, Patient Name, Age, Gender, and Medical Record Number (MRN).")
    add_bullet("Shift Metadata: ", "Outgoing shift type, outgoing primary Registered Nurse (RN), incoming primary RN, and timestamp.")
    add_bullet("Longitudinal Shift Vitals & Trends: ", "Shift vital sign telemetry, oxygen titration levels, pain scores, and calculated EWS (Early Warning Score) delta.")
    add_bullet("eMAR Medication Administration Data: ", "Medications administered during the shift, pending doses, overdue doses, and PRN (as-needed) analgesic records.")
    add_bullet("High-Alert Medication Orders: ", "Active infusions of high-alert medications (Insulin, Heparin, Vancomycin, Narcotics, Vasopressors, Inotropes).")
    add_bullet("Nursing Clinical Assessments: ", "Neurological status, intake/output fluid balance, surgical drain outputs, IV (Intravenous) line insertion site status, wound dressing condition, and fall risk score.")
    add_bullet("Active Doctor Orders: ", "New stat physician orders, pending urgent lab investigations, and scheduled imaging transfers.")
    
    add_h3("Outputs Produced by the Agent:")
    add_bullet("Structured Bedside SBAR Handover Card: ", "Clinically synthesized Situation, Background, Assessment, and Recommendation summary tailored for the ward RN.")
    add_bullet("High-Alert Drug Safety Flags: ", "Highlighted callout banners for narrow therapeutic index and high-alert drug infusions requiring joint double-check.")
    add_bullet("Clinical Risk & Deterioration Alerts: ", "Early detection warnings for rising EWS (Early Warning Score), abnormal fluid balance deficits, or pending urgent lab follow-ups.")
    add_bullet("Interactive Shift Checklist: ", "Actionable tasks checklist for the incoming nurse (e.g., scheduled antibiotic infusions, dressing changes, fasting status).")
    add_bullet("Persisted Bedside Audit Record: ", "Signed handover summary written to ward_sbar_handovers upon dual RN electronic acknowledgment.")
    
    add_h3("Technical & Operational Details:")
    add_bullet("How Staff Use It: ", "Ward Tablet / Nursing Workspace Bed Board. Both outgoing and incoming nurses review the digital SBAR card and tap 'Acknowledge & Accept'.")
    add_bullet("Delivery Mode & Primary Users: ", "Ward Bedside Workflow Automation (Nursing Tablet). Primary users: Ward Staff Nurses and Charge Nurses.")
    add_bullet("LLM Models: ", "Primary: openai/gpt-oss-120b on Groq LPU (NURSING_AGENT_MODEL). Fallback: Gemini 3.5 Flash-Lite. Temp: 0.15. p50 Latency: ~1.85s.")
    add_bullet("Tools & Integrations: ", "EMR API (Vitals, Clinical Notes), Pharmacy API (High-Alert / eMAR), Document Generator (SBAR Cards).")
    add_bullet("Knowledge & SOPs: ", "Medication Safety — High-Alert Drugs SOP v4.0, Inpatient Ward Administration Guidelines v3.2.")
    add_bullet("Quality Benchmarks: ", "Studio Evaluation: 94.8% clinical accuracy, 97.5% groundedness, 0.3% hallucination rate, ~1.85s response time.", space_after=12)

    # -------------------------------------------------------------
    # AG-04
    # -------------------------------------------------------------
    add_h2("AG-04 · Employee Self-Service & HR Assistant")
    add_p("Human Resources (HR) Operations · Internal Staff Support Tier · Published · v1.7.0", bold_prefix="Domain & Governance: ", italic=True)
    
    add_p(
        "Acts as an intelligent 24/7 conversational assistant for hospital staff (physicians, nurses, allied health workers, and admin) "
        "to instantly resolve queries regarding duty rosters, shift schedules, leave balances (Casual, Sick, Earned, Comp-Off), overtime policies, "
        "and automatically file leave and comp-off applications for supervisory approval.",
        bold_prefix="Purpose: "
    )
    
    add_p(
        "Eliminates friction and delays in HR administrative tasks. Nurses and hospital staff get instant answers regarding upcoming shifts "
        "and leave entitlements without manual phone calls, paperwork, or WhatsApp messages to busy charge nurses and HR managers.",
        bold_prefix="Why It Is Useful: "
    )
    
    add_h3("Inputs Given to the Agent:")
    add_bullet("User Identity & Role: ", "Employee ID, Staff Name, Department, Clinical Designation (Consultant, Staff RN, Lab Tech, Administrative Staff).")
    add_bullet("Conversational User Prompt: ", "Natural language user text input (e.g., 'What shift am I assigned to tomorrow?', 'Check my available Casual Leave', 'Apply comp-off for Sunday duty').")
    add_bullet("Staff Duty Rosters: ", "Live departmental duty scheduling records from PostgreSQL staff_rosters table.")
    add_bullet("Employee Leave Ledger: ", "Current accrued balances across Casual Leave (CL), Sick Leave (SL), Earned Leave (EL), and Compensatory Off (Comp-Off) from employee_leave_balances table.")
    add_bullet("Leave Request Parameters: ", "Selected leave category, start date, end date, total days, shift coverage substitute, and reason.")
    add_bullet("Hospital HR Policies: ", "Official documentation including HR Leave & Attendance Policy v5.0, Payroll Rules, and Shift Allowance Guidelines.")
    
    add_h3("Outputs Produced by the Agent:")
    add_bullet("Conversational Roster & Schedule Responses: ", "Immediate natural-language clarification of upcoming shift assignments, timing, and ward allocations.")
    add_bullet("Visual Leave Balance Cards: ", "Clean breakdown of total entitlement, utilized days, and remaining balance per leave category.")
    add_bullet("Validated Leave & Comp-Off Submissions: ", "Structured leave application records validated against business rules and inserted into employee_leave_requests.")
    add_bullet("Automated Supervisory Notifications: ", "Instant routing of leave requests to the respective Department Head or Nurse Supervisor for digital approval.")
    add_bullet("HR Policy Explanations: ", "Contextual answers regarding night shift allowances, public holiday compensations, and maternity/medical leave rules.")
    
    add_h3("Technical & Operational Details:")
    add_bullet("How Staff Use It: ", "Interactive Chat Widget / Employee Portal. Staff interact via natural language or quick-action chips ('My Shift Tomorrow', 'Apply Comp-Off').")
    add_bullet("Delivery Mode & Primary Users: ", "Conversational Internal Staff Assistant (Web & Mobile). Primary users: Doctors, Nurses, Technicians, Hospital Staff.")
    add_bullet("LLM Models: ", "Runtime: Groq Chat Completions (openai/gpt-oss-120b or Meta-Llama-3.3-70B). Secondary fallback: Gemini via LLM_API_KEY. Temp: ~0.20.")
    add_bullet("Tools & Integrations: ", "PostgreSQL Roster Engine, Leave Ledger API, Supervisor Notification Router, HR Policy Retrieval Engine.")
    add_bullet("Knowledge & SOPs: ", "HR Leave & Attendance Policy v5.0, Payroll & Shift Allowance FAQ v2.1, Nursing Staffing SOP v3.2.")
    add_bullet("Quality Benchmarks: ", "Internal Profile Evaluation: 98.5% intent accuracy, 99.2% policy groundedness, 0% clinical permission risk.", space_after=12)

    # -------------------------------------------------------------
    # AG-14
    # -------------------------------------------------------------
    add_h2("AG-14 · Executive Hospital Operations & Analytics Agent")
    add_p("Management & Executive Leadership · Strategic Analytics Tier · Published · v1.7.0", bold_prefix="Domain & Governance: ", italic=True)
    
    add_p(
        "Transforms live operational, clinical, and financial database facts into real-time executive KPI (Key Performance Indicator) "
        "dashboards and structured natural-language daily briefings covering bed occupancy, emergency inflow, billing realization, "
        "insurance claim conversion, and clinical department workloads.",
        bold_prefix="Purpose: "
    )
    
    add_p(
        "Eliminates dependency on delayed weekly manual spreadsheet reports. Hospital leadership (CEO, CFO, Medical Director) can monitor "
        "live operational health, track revenue velocity, ask complex analytical questions in plain English, and instantly uncover departmental bottlenecks.",
        bold_prefix="Why It Is Useful: "
    )
    
    add_h3("Inputs Given to the Agent:")
    add_bullet("Inpatient Census & Flow Records: ", "Live admission inputs, discharges, bed occupancy counts, and ED (Emergency Department) triage logs.")
    add_bullet("Financial & Billing Data: ", "Total gross billing, discounts, insurance deductions, net collected revenue, and outstanding accounts receivable.")
    add_bullet("Insurance & TPA Metrics: ", "Insurance claims submitted, initial approvals, shortfalls, rejections, and settlement turnaround times.")
    add_bullet("Departmental & Clinical Workloads: ", "Consultation volumes, surgical theatre utilization rates, diagnostic lab test volumes across specialties (Cardiology, Orthopedics, Oncology, etc.).")
    add_bullet("Executive Natural-Language Queries: ", "Ad-hoc management questions (e.g., 'Compare Cardiology vs Orthopedics admission pace and revenue this quarter').")
    add_bullet("Historical Target Benchmarks: ", "Hospital budgetary revenue targets, NABH average length of stay (ALOS) targets, and occupancy thresholds.")
    
    add_h3("Outputs Produced by the Agent:")
    add_bullet("Executive KPI Dashboard Cards: ", "Real-time metrics: Overall Bed Occupancy %, Emergency Flow Rate, Cash vs Insurance Realization %, and Average Length of Stay.")
    add_bullet("Executive Natural-Language Briefings: ", "AI-synthesized morning and evening operational summaries detailing key achievements, risks, and performance highlights.")
    add_bullet("Specialty Comparative Breakdown: ", "Comparative charts and revenue/occupancy tables dissecting performance across medical specialties and doctor panels.")
    add_bullet("Operational Bottleneck Alerts: ", "Early detection warnings for prolonged discharge turnaround times, delayed TPA pre-authorization cycles, or high emergency boarding.")
    add_bullet("De-Identified Data Stream: ", "Strictly aggregated analytical output maintaining complete Protected Health Information (PHI) boundary integrity.")
    
    add_h3("Technical & Operational Details:")
    add_bullet("How Staff Use It: ", "AnalyticsView / Agent Studio. Real-time data ingested via GET /api/v1/gold/live-analytics. Leadership reviews cards and natural-language briefings.")
    add_bullet("Delivery Mode & Primary Users: ", "Hybrid Executive Dashboard & Natural-Language Analytical Interface. Primary users: CEO, CFO, Medical Director, Operations Heads.")
    add_bullet("LLM Models: ", "Narrative & query layer: openai/gpt-oss-120b on Groq. Fallback: Gemini 3.5 Flash-Lite. Temp: ~0.10. Note: Numerical KPIs are deterministic SQL.")
    add_bullet("Tools & Integrations: ", "HMS Census & KPI Engine, Billing Revenue / TAT Analytics Engine, Executive Document Briefing Generator.")
    add_bullet("Knowledge & SOPs: ", "Inpatient Discharge SOP v3.1, Hospital Tariff Schedule FY26-27, Executive Governance Charter.")
    add_bullet("Quality Benchmarks: ", "Studio Evaluation EV-714: 93.8% accuracy, 97.0% groundedness, 0.3% hallucination rate.", space_after=12)

    # -------------------------------------------------------------
    # AG-15
    # -------------------------------------------------------------
    add_h2("AG-15 · Inpatient Bed Census & Ward Capacity Forecasting Agent")
    add_p("Hospital Operations & Bed Management · Predictive Capacity Tier · Published · v1.0.6", bold_prefix="Domain & Governance: ", italic=True)
    
    add_p(
        "Leverages predictive machine learning and statistical time-series forecasting to project 7-day forward inpatient census, "
        "ward-level bed occupancy, and capacity headroom across General Wards, Private Suites, CCU (Coronary Care Unit), and ICU (Intensive Care Unit), "
        "alerting hospital leadership to impending capacity crunches before they occur.",
        bold_prefix="Purpose: "
    )
    
    add_p(
        "Enables proactive capacity management. For example, on Monday morning, the COO and bed management team can forecast a critical "
        "ICU surge heading above 85% by Thursday, allowing them to proactively transition stable patients to step-down units or arrange "
        "staffing adjustments, thereby preventing emergency boarding and elective surgery cancellations.",
        bold_prefix="Why It Is Useful: "
    )
    
    add_h3("Inputs Given to the Agent:")
    add_bullet("Real-Time Bed Occupancy Ledger: ", "Live bed counts, occupied beds, vacant beds, and beds out of service across all hospital wards and ICUs.")
    add_bullet("Inpatient Admission & Discharge Velocity: ", "Historical admissions and discharge pacing profiles, segmented by day of week and clinical specialty.")
    add_bullet("Scheduled Elective Admissions: ", "Pre-booked surgical admissions and procedural scheduling records from the hospital OT registry.")
    add_bullet("Emergency Department Triage Velocity: ", "Live ED arrival velocity, admission conversion ratios, and emergency bed request queues.")
    add_bullet("Ward Mix & Patient Acuity Data: ", "Clinical acuity stratification (ICU, Step-Down, Telemetry, General Ward) and anticipated length-of-stay distributions.")
    
    add_h3("Outputs Produced by the Agent:")
    add_bullet("7-Day Predictive Census Trajectory: ", "Day-by-day projected census (T+0 through T+6) with statistical headroom confidence intervals.")
    add_bullet("Ward-Level Capacity Warnings: ", "Proactive warning badges highlighting specific wards projected to exceed safety thresholds (e.g., 'ICU capacity heading to 88% on Thursday').")
    add_bullet("Actionable Bed Reallocation Guidance: ", "Strategic suggestions for step-down transfers and surge bed activations to maintain critical care buffers.")
    add_bullet("Nursing Shift Requirement Forecast: ", "Projected nurse-to-patient staffing ratio requirements tailored to anticipated ward occupancy spikes.")
    add_bullet("Predictive Sentinel Briefing: ", "Natural-language narrative overview explaining underlying admission pacing drivers and seasonal surge factors.")
    
    add_h3("Technical & Operational Details:")
    add_bullet("How Staff Use It: ", "Live Forecasting Dashboard. Ingests data via GET /api/v1/gold/live-forecasting. Supervisors use projections to plan ward allocations.")
    add_bullet("Delivery Mode & Primary Users: ", "Background Capacity Sentinel & Bed Operations Dashboard. Primary users: COO, Bed Managers, Nursing Supervisors.")
    add_bullet("LLM Models: ", "Predictive Engine: LightGBM + Prophet Inpatient Census Predictor v2.4 (deterministic math). Narrative: openai/gpt-oss-120b on Groq. Temp: ~0.10.")
    add_bullet("Tools & Integrations: ", "HMS Bed Demand & ER Flow API, Nursing Shift Scheduling Integration, Capacity Alert Dispatcher.")
    add_bullet("Knowledge & SOPs: ", "Inpatient Discharge SOP v3.1, Critical Care Triage & Bed Reallocation Guidelines.")
    add_bullet("Quality Benchmarks: ", "Studio Evaluation EV-715: 94.0% predictive accuracy; live validation reports strong correlation (R² ~0.942).", space_after=12)

    # -------------------------------------------------------------
    # AG-07
    # -------------------------------------------------------------
    add_h2("AG-07 · Cashless Insurance Pre-Authorization Dossier Agent")
    add_p("Insurance Desk · High Financial & Operational Tier · Published · v2.1.0", bold_prefix="Domain & Governance: ", italic=True)
    
    add_p(
        "Upon inpatient admission or procedural scheduling of an insured patient, this agent automatically synthesizes a complete, "
        "fully compliant cashless pre-authorization dossier—including clinical history, ICD-10 diagnostic coding, itemized tariff estimates, "
        "document checklist verification, and AI denial-risk scoring—ready for one-click submission to the TPA (Third-Party Administrator) "
        "following human officer review.",
        bold_prefix="Purpose: "
    )
    
    add_p(
        "Reduces manual pre-authorization preparation time from 45 minutes of hunting across paper charts and billing ledgers down to "
        "~20 seconds of officer validation. Dramatically enhances first-pass approval rates from insurers, minimizes administrative rejections, "
        "and eliminates discharge delays caused by pending insurance approvals.",
        bold_prefix="Why It Is Useful: "
    )
    
    add_h3("Inputs Given to the Agent:")
    add_bullet("Patient Demographic & Insurance Identity: ", "Patient Name, Age, Gender, Policy Number, Health Card Number, TPA Name, Insurer, and Sum Insured.")
    add_bullet("Clinical Admission & Physician Advice: ", "Treating Doctor clinical notes, provisional diagnosis, presenting complaints, and proposed treatment plan.")
    add_bullet("Diagnostic & Procedure Coding: ", "ICD-10 (International Classification of Diseases) diagnostic codes and PCS (Procedure Coding System) surgical codes.")
    add_bullet("Itemized Hospital Cost Estimate: ", "Room rent charges, OT fees, surgeon fees, planned consumable/implant estimates, and diagnostic investigation costs.")
    add_bullet("Supporting Clinical Documentation: ", "Attached diagnostic reports, physician consultation slips, pre-admission investigation results, and identity proofs.")
    add_bullet("Regulatory & TPA Rule Sets: ", "IRDAI (Insurance Regulatory and Development Authority of India) Preauth Guidelines 2024 and specific TPA tariff/exclusion rules.")
    
    add_h3("Outputs Produced by the Agent:")
    add_bullet("Complete Cashless Preauth Dossier: ", "Standardized, structured insurance pre-authorization packet populated with validated clinical and financial fields.")
    add_bullet("Two-Stage AI Denial-Risk Assessment: ", "Calculated Denial-Risk Badge (Low, Moderate, High Risk) detailing specific risk factors (e.g., missing conservative therapy history).")
    add_bullet("Missing Documentation Checklist: ", "Automated alert flags for any missing mandatory documents (e.g., missing pre-admission investigation report or signed estimate).")
    add_bullet("Standardized Clinical Justification Draft: ", "AI-formulated formal medical necessity justification letter addressed to the TPA medical officer.")
    add_bullet("1-Click Submission Trigger: ", "Actionable interface enabling the insurance officer to review, modify, and transmit the packet via the TPA portal gateway.")
    
    add_h3("Technical & Operational Details:")
    add_bullet("How Staff Use It: ", "Insurance Workspace / Kanban Board + Preauth Dossier Drawer. Officer reviews checklist and clicks 'Submit Preauth Packet to TPA'. Zero autonomous send.")
    add_bullet("Delivery Mode & Primary Users: ", "Workflow Automation System. Primary users: Insurance Desk Executives (e.g., R. Sundar, L. Fathima), Billing Officers.")
    add_bullet("LLM Models: ", "Primary: openai/gpt-oss-120b on Groq LPU + Policy & Clinical Validation Engine (preauth-denial v0.9). Fallback: Structured dossier builder. Temp: Low.")
    add_bullet("Tools & Integrations: ", "EMR API, Insurance/TPA Gateway API, Hospital Billing & Tariff API, Document Dossier Generator.")
    add_bullet("Knowledge & SOPs: ", "IRDAI Preauth Guidelines 2024, TPA Tariff & Denial Code Catalogs, Star Health / ICICI Lombard Preauth Checklists, SOP v2.4.")
    add_bullet("Quality Benchmarks: ", "Turnaround time: 20 seconds vs 45 min manual; 99.1% profile accuracy; 96.4% first-pass insurer approval rate.", space_after=12)

    # -------------------------------------------------------------
    # AG-20
    # -------------------------------------------------------------
    add_h2("AG-20 · Insurance Claim Shortfall & Denial Appeal Agent")
    add_p("Insurance Desk & RCM · Medium Financial Recovery Tier · Published · v1.0.2", bold_prefix="Domain & Governance: ", italic=True)
    
    add_p(
        "When an insurance company or TPA (Third-Party Administrator) rejects an inpatient claim or imposes an unfair shortfall deduction, "
        "this agent automatically analyzes the denial code, matches it against patient EMR (Electronic Medical Record) and OT (Operation Theatre) clinical proof, "
        "and drafts a formal, legally grounded reconsideration appeal letter with an attached evidence checklist for one-click human submission.",
        bold_prefix="Purpose: "
    )
    
    add_p(
        "Prevents severe hospital revenue leakage caused by routine insurance shortfalls (e.g., unapproved consumable deductions, room rent capping, "
        "alleged lack of conservative management proof). The insurance officer instantly receives an evidence-backed appeal packet, enabling the "
        "hospital to recover significant funds (e.g., ₹80,000 spinal-fusion shortfall recovery) that would otherwise be written off.",
        bold_prefix="Why It Is Useful: "
    )
    
    add_h3("Inputs Given to the Agent:")
    add_bullet("TPA Claim Settlement Voucher: ", "Official denial notice, settlement letter, rejected line items, and disallowed claim amounts.")
    add_bullet("Standardized Claim Denial Codes: ", "TPA denial/shortfall reason codes (e.g., Code 204: Missing Conservative Management; Code 402: Room Rent Disallowance; Code 108: Non-Medical Consumables).")
    add_bullet("Original Hospital Billing Ledger: ", "Comprehensive itemized hospital invoice, pharmacy dispensary records, and tariff schedule.")
    add_bullet("Clinical Inpatient Medical Records: ", "EMR clinical notes, OT operative notes, anesthesia logs, nursing administration charts, and physician daily progress entries.")
    add_bullet("Regulatory Appeals Framework: ", "IRDAI Cashless Settlement Guidelines 2024, Master Circular 2020 on Claim Repudiations, and Hospital Tariff Schedule FY26-27.")
    
    add_h3("Outputs Produced by the Agent:")
    add_bullet("Denial Root-Cause & Disputed Sum Analysis: ", "Detailed breakdown of the exact financial dispute, disallowed amounts, and regulatory grounds for contestation.")
    add_bullet("Matched Clinical Evidence Package: ", "Curated clinical proof extracted from EMR (e.g., highlighted OT notes proving surgical necessity, failed conservative physiotherapy records).")
    add_bullet("Formal Reconsideration Appeal Letter Draft: ", "Comprehensive, professional legal-medical appeal letter addressed to the TPA Grievance / Claims Lead citing IRDAI clauses.")
    add_bullet("Claim Appeal Database Record: ", "Structured audit entry persisted to claim_appeals database table tracking appeal lifecycle status.")
    add_bullet("1-Click Officer Submission Dossier: ", "Insurance desk review card with one-click approval button to transmit the appeal packet to the insurer.")
    
    add_h3("Technical & Operational Details:")
    add_bullet("How Staff Use It: ", "Insurance Claims Desk. Officer opens the 'Shortfall Detected' card, reviews auto-compiled evidence and letter, and clicks 'Submit Appeal to TPA'.")
    add_bullet("Delivery Mode & Primary Users: ", "Workflow Automation System. Primary users: Insurance Desk Executives, Revenue Cycle Management (RCM) Leads.")
    add_bullet("LLM Models: ", "Live appeal generation utilizes grounded clinical evidence templates keyed by denial codes + openai/gpt-oss-120b on Groq for contextual narrative synthesis. Temp: ~0.20.")
    add_bullet("Tools & Integrations: ", "TPA Rejection Parser, EMR Clinical Retrieval Engine, Appeal Dossier Assembler.")
    add_bullet("Knowledge & SOPs: ", "Insurance Preauth SOP v2.4, IRDAI Cashless Settlement Guidelines 2024, IRDAI Master Circular 2020, Tariff FY26-27.")
    add_bullet("Quality Benchmarks: ", "Appeal Success Rate: ~91.0% first-pass appeal recovery rate across evaluated shortfall dispute cases.", space_after=12)

    # -------------------------------------------------------------
    # AG-08
    # -------------------------------------------------------------
    add_h2("AG-08 · Billing Transparency & Estimate Variance Explainer Agent")
    add_p("Hospital Finance & Billing Desk · High Financial & Patient Experience Tier · Published · v2.0.0", bold_prefix="Domain & Governance: ", italic=True)
    
    add_p(
        "Continuously audits the running inpatient hospital bill against the initial pre-admission cost estimate. If financial variance "
        "exceeds ~10%, the agent correlates technical billing line items with EMR (Electronic Medical Record) and OT (Operation Theatre) clinical notes "
        "to generate plain-language bilingual (English and Tamil) explanations of why additional procedures, implants, or ICU days occurred.",
        bold_prefix="Purpose: "
    )
    
    add_p(
        "Eliminates contentious billing disputes at the discharge billing counter. Families often get frustrated by cryptic technical line items "
        "(e.g., 'NC Balloon Catheter 2.5x15mm' or 'Extra ICU monitoring'). By providing clear, bilingual, clinically justified explanations on screen, "
        "cashiers resolve questions immediately, accelerate cash collection, and protect patient trust without altering official tariffs.",
        bold_prefix="Why It Is Useful: "
    )
    
    add_h3("Inputs Given to the Agent:")
    add_bullet("Pre-Admission Cost Estimate Ledger: ", "Initial signed financial estimate sheet including anticipated length of stay, procedure cost, and room category.")
    add_bullet("Real-Time Running Hospital Bill: ", "Live itemized billing ledger containing room charges, pharmacy dispensations, OT consumable charges, lab fees, and doctor visits.")
    add_bullet("Financial Variance Trigger: ", "Automated threshold flag triggered when actual running charges exceed the initial estimate by >10%.")
    add_bullet("Surgical & Operative Clinical Records: ", "OT operative summary, surgeon clinical progress notes, and intra-operative complication logs explaining procedural changes.")
    add_bullet("Multilingual Medical Lexicon: ", "Curated clinical dictionary translating complex medical terms and surgical hardware into patient-friendly English and Tamil descriptions.")
    add_bullet("Hospital Tariff Master: ", "Hospital Tariff Schedule FY26-27 defining standardized line item rates and consumable pricing guidelines.")
    
    add_h3("Outputs Produced by the Agent:")
    add_bullet("Estimate vs Actual Variance Breakdown Card: ", "Clear visual comparison highlighting exactly which department or category (e.g., OT consumables, extra ICU day) caused the variance.")
    add_bullet("Plain-Language Bilingual Explanations: ", "Easy-to-understand explanations in English and Tamil clarifying technical medical items (e.g., explaining why a second coronary balloon catheter was necessary during angioplasty).")
    add_bullet("Clinical Necessity Proof Badges: ", "Direct clickable references linking each variance line item to the specific treating doctor's OT note or clinical order.")
    add_bullet("Investigate Clinical Necessity Drawer: ", "Interactive drill-down interface for the cashier and patient family detailing timeline of events during the hospital stay.")
    add_bullet("Printable Bilingual Bill Explanation Slip: ", "Optional patient-facing summary slip ready for cashier electronic sign-off and physical counter handout.")
    
    add_h3("Technical & Operational Details:")
    add_bullet("How Staff Use It: ", "Billing Desk / Discharge Counter / Financial Revenue screens. Orange 'Estimate Variance' badge alerts cashier. Cashier presents the explanation card.")
    add_bullet("Delivery Mode & Primary Users: ", "Embedded Billing Desk Card & Financial Operations Interface. Primary users: Billing Executives, Cashiers, Ward Administrators, Patients/Families.")
    add_bullet("LLM Models: ", "Primary: openai/gpt-oss-120b on Groq LPU (BILLING_AGENT_MODEL / NURSING_AGENT_MODEL). Fallback: Gemini 3.5 Flash-Lite or rule-based clinical synthesis. Temp: ~0.20.")
    add_bullet("Tools & Integrations: ", "Hospital Billing Desk API, Pre-Admission Estimate Ledger, EMR & OT Clinical Notes Retrieval, Bilingual Explainer Generator.")
    add_bullet("Knowledge & SOPs: ", "Hospital Tariff Schedule FY26-27, Consumable & Implant Medical Nomenclature, Tamil Medical Lexicon.")
    add_bullet("Quality Benchmarks: ", "Studio Evaluation EV-708: 96.4% explanation accuracy, 98.2% groundedness, 0.2% hallucination rate, ~1.4s response latency.", space_after=12)

    # -------------------------------------------------------------
    # SECTION 4: ACRONYM GLOSSARY & REFERENCE TABLE
    # -------------------------------------------------------------
    add_h1("4. Standard Healthcare & Technical Acronym Reference")
    
    add_p(
        "For complete clarity and standard cross-departmental documentation, the following standardized abbreviations and expansions "
        "are utilized across all eight AI agents:"
    )
    
    acronyms = [
        ("EMR", "Electronic Medical Record (Digital clinical record of patient care within the hospital)"),
        ("EHR", "Electronic Health Record (Comprehensive longitudinal health record across healthcare systems)"),
        ("NABH", "National Accreditation Board for Hospitals & Healthcare Providers (Apex healthcare quality accreditation body in India)"),
        ("SBAR", "Situation, Background, Assessment, Recommendation (Standardized clinical handover communication framework)"),
        ("eMAR", "Electronic Medication Administration Record (Digital nursing log of administered medications and dosages)"),
        ("EWS", "Early Warning Score (Physiological scoring system used to quickly determine patient clinical deterioration)"),
        ("RN", "Registered Nurse (Licensed professional nurse responsible for inpatient bedside care)"),
        ("TPA", "Third-Party Administrator (Organization processing cashless insurance claims on behalf of insurance companies)"),
        ("IRDAI", "Insurance Regulatory and Development Authority of India (Statutory regulatory body overseeing insurance in India)"),
        ("ICD-10", "International Classification of Diseases, 10th Revision (Standard diagnostic coding system maintained by WHO)"),
        ("OT", "Operation Theatre (Sterile surgical suite where operative procedures are performed)"),
        ("ICU", "Intensive Care Unit (Specialized critical care facility for critically ill patients)"),
        ("CCU", "Coronary Care Unit (Specialized intensive care unit for cardiac patients)"),
        ("ED / ER", "Emergency Department / Emergency Room (Hospital department providing immediate acute medical care)"),
        ("COO", "Chief Operating Officer (Executive responsible for day-to-day hospital operations and capacity)"),
        ("CEO", "Chief Executive Officer (Chief executive responsible for overall organizational management)"),
        ("CFO", "Chief Financial Officer (Executive managing hospital finances, billing, and revenue cycle)"),
        ("HR", "Human Resources (Department overseeing workforce staffing, rosters, and employee welfare)"),
        ("KPI", "Key Performance Indicator (Measurable value evaluating organizational operational and financial success)"),
        ("LLM", "Large Language Model (Deep learning algorithm trained on massive datasets to understand and generate text)"),
        ("LPU", "Language Processing Unit (High-speed specialized hardware inference engine developed by Groq)"),
        ("LIS", "Laboratory Information System (Software system that manages hospital diagnostic lab data)"),
        ("HMS", "Hospital Management System (Integrated software platform managing hospital clinical, administrative, and financial operations)"),
        ("SOP", "Standard Operating Procedure (Established stepwise guidelines for clinical and operational hospital workflows)"),
        ("TAT", "Turnaround Time (Time taken from the initiation to the completion of a specific workflow process)"),
        ("PHI", "Protected Health Information (Sensitive, individually identifiable patient health data protected under privacy laws)"),
        ("SQL", "Structured Query Language (Standard domain-specific programming language for querying relational databases)"),
        ("API", "Application Programming Interface (Software intermediary allowing distinct computer systems to communicate)")
    ]
    
    for acr, expansion in acronyms:
        add_bullet(f"{acr}: ", expansion, space_after=2)

    # Save document
    primary_path = r'e:\Bosco-projects\POC\Health-care\code\health-care\docs\Eight_Hospital_AI_Agents_Details.docx'
    updated_path = r'e:\Bosco-projects\POC\Health-care\code\health-care\docs\Eight_Hospital_AI_Agents_Details_Updated.docx'
    
    # Save to updated path first
    doc.save(updated_path)
    print(f"Document successfully created and saved to: {updated_path}")
    
    # Try saving to primary path if unlocked
    try:
        doc.save(primary_path)
        print(f"Document also successfully updated at: {primary_path}")
    except PermissionError:
        print(f"Notice: '{primary_path}' is currently open in Microsoft Word. Saved to '{updated_path}'. Close Word to overwrite the original.")

if __name__ == '__main__':
    create_document()
